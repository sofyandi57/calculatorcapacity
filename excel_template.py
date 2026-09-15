"""
excel_template.py
Membangun file Excel "Kalkulator Infrastruktur" dengan struktur & rumus hidup
(bukan angka statis) — meniru template standar: sheet Dashboard, Input - Infra
Utama, Input - Infra Pendukung, Input - Storage Capacity, Input - Network
Capacity, dan Output & Gap Analysis. Sel kuning = boleh diedit manual, sel abu
= formula (jangan diedit). Dipakai sebagai default template saat download Excel.
"""

from __future__ import annotations

from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(name="Arial", italic=True, size=9, color="595959")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")

MIN_PLATFORM_ROWS = 8
MIN_PRODUK_ROWS = 10
MIN_UTAMA_ROWS = 20
MIN_PENDUKUNG_ROWS = 10
MIN_STORAGE_ROWS = 15
MIN_NETWORK_ROWS = 10
MIN_MEMBER_ROWS = 8


def _title(ws, text: str, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT


def _subtitle(ws, row: int, text: str, ncols: int):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = NOTE_FONT
    c.alignment = WRAP_LEFT


def _header_row(ws, row: int, headers: List[str], col_widths: List[float] | None = None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = WRAP_CENTER
        if col_widths:
            ws.column_dimensions[get_column_letter(i)].width = col_widths[i - 1]
    ws.row_dimensions[row].height = 30


def _set(ws, row: int, col: int, value, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:
        c.fill = fill
    return c


def build_template_workbook(DATA: Dict[str, Any]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    platforms = DATA.get("platforms", [])
    infra_utama = DATA.get("infra_utama", [])
    pendukung = DATA.get("pendukung", [])
    storage = DATA.get("storage", [])
    network = DATA.get("network", [])
    network_members = DATA.get("network_members", [])

    n_platform = max(MIN_PLATFORM_ROWS, len(platforms))
    n_produk = max(MIN_PRODUK_ROWS, len({(r.get("platform"), r.get("produk")) for r in infra_utama}))
    n_utama = max(MIN_UTAMA_ROWS, len(infra_utama))
    n_pendukung = max(MIN_PENDUKUNG_ROWS, len(pendukung))
    n_storage = max(MIN_STORAGE_ROWS, len(storage))
    n_network = max(MIN_NETWORK_ROWS, len(network))
    n_member = max(MIN_MEMBER_ROWS, len(network_members))

    UTAMA_FIRST_ROW = 7
    UTAMA_LAST_ROW = UTAMA_FIRST_ROW + n_utama - 1
    PENDUKUNG_FIRST_ROW = 5
    PENDUKUNG_LAST_ROW = PENDUKUNG_FIRST_ROW + n_pendukung - 1
    STORAGE_FIRST_ROW = 6
    STORAGE_LAST_ROW = STORAGE_FIRST_ROW + n_storage - 1
    NETWORK_FIRST_ROW = 6
    NETWORK_LAST_ROW = NETWORK_FIRST_ROW + n_network - 1
    MEMBER_FIRST_ROW = NETWORK_LAST_ROW + 5
    MEMBER_LAST_ROW = MEMBER_FIRST_ROW + n_member - 1

    # =====================================================================
    # SHEET: Input - Infra Utama
    # =====================================================================
    ws = wb.create_sheet("Input - Infra Utama")
    _title(ws, "Kalkulator Infrastruktur — INFRA UTAMA (basis Current & Future Transaction Handling / TPS)", 26)
    _subtitle(ws, 2, "Setiap Produk dipecah menjadi 2 baris: Sistem DB dan Sistem Engine (karakteristik resource berbeda).", 26)

    headers = ["No", "Platform", "Produk", "Sistem", "Kategori\nInfrastruktur",
               "Current TPS", "Future TPS", "CPU/TPS\n(vCore)", "Memory/TPS\n(GB)", "Storage/TPS\n(GB)",
               "Req @Current CPU", "Req @Current Mem", "Req @Current Storage",
               "Req @Future CPU", "Req @Future Mem", "Req @Future Storage",
               "Buffer /\nHeadroom (%)", "HA\nMultiplier",
               "FINAL CPU", "FINAL Memory", "FINAL Storage",
               "Kuota CPU\nProduk (vCore)", "Kuota Memory\nProduk (GB)",
               "Total FINAL CPU\nper Produk (rollup)", "Total FINAL Memory\nper Produk (rollup)",
               "Status CPU\n(vs Kuota Produk)", "Status Memory\n(vs Kuota Produk)"]
    _header_row(ws, 4, headers, col_widths=[5, 16, 22, 9, 10, 10, 10, 9, 9, 9, 10, 10, 11, 10, 10, 11, 9, 8, 10, 10, 10, 10, 10, 11, 11, 11, 11])

    for i in range(n_utama):
        row = UTAMA_FIRST_ROW + i
        data = infra_utama[i] if i < len(infra_utama) else {}
        _set(ws, row, 1, i + 1)
        _set(ws, row, 2, data.get("platform", ""), INPUT_FILL)
        _set(ws, row, 3, data.get("produk", ""), INPUT_FILL)
        _set(ws, row, 4, data.get("sistem", ""), INPUT_FILL)
        _set(ws, row, 5, data.get("kategori", "Utama"), INPUT_FILL)
        _set(ws, row, 6, data.get("current_tps"), INPUT_FILL)
        _set(ws, row, 7, data.get("future_tps"), INPUT_FILL)
        _set(ws, row, 8, data.get("rasio_cpu"), INPUT_FILL)
        _set(ws, row, 9, data.get("rasio_mem"), INPUT_FILL)
        _set(ws, row, 10, data.get("rasio_storage"), INPUT_FILL)
        _set(ws, row, 11, f"=F{row}*H{row}", FORMULA_FILL)
        _set(ws, row, 12, f"=F{row}*I{row}", FORMULA_FILL)
        _set(ws, row, 13, f"=F{row}*J{row}", FORMULA_FILL)
        _set(ws, row, 14, f"=G{row}*H{row}", FORMULA_FILL)
        _set(ws, row, 15, f"=G{row}*I{row}", FORMULA_FILL)
        _set(ws, row, 16, f"=G{row}*J{row}", FORMULA_FILL)
        _set(ws, row, 17, (data.get("buffer_pct", 20.0) / 100.0) if data else 0.2, INPUT_FILL)
        _set(ws, row, 18, data.get("ha_multiplier", 1.0) if data else 1, INPUT_FILL)
        _set(ws, row, 19, f"=N{row}*(1+Q{row})*R{row}", FORMULA_FILL)
        _set(ws, row, 20, f"=O{row}*(1+Q{row})*R{row}", FORMULA_FILL)
        _set(ws, row, 21, f"=P{row}*(1+Q{row})*R{row}", FORMULA_FILL)
        _set(ws, row, 22, data.get("kuota_cpu"), INPUT_FILL)
        _set(ws, row, 23, data.get("kuota_mem"), INPUT_FILL)
        _set(ws, row, 24, f"=SUMIF($C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},$C{row},$S${UTAMA_FIRST_ROW}:$S${UTAMA_LAST_ROW})", FORMULA_FILL)
        _set(ws, row, 25, f"=SUMIF($C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},$C{row},$T${UTAMA_FIRST_ROW}:$T${UTAMA_LAST_ROW})", FORMULA_FILL)
        _set(ws, row, 26, f'=IF(V{row}="","",IF(X{row}/V{row}>1,"OVER KUOTA",IF(X{row}/V{row}>0.7,"WARNING","OK")))', FORMULA_FILL)
        _set(ws, row, 27, f'=IF(W{row}="","",IF(Y{row}/W{row}>1,"OVER KUOTA",IF(Y{row}/W{row}>0.7,"WARNING","OK")))', FORMULA_FILL)

    total_row = UTAMA_LAST_ROW + 1
    _set(ws, total_row, 1, "TOTAL INFRA UTAMA")
    for col, letter in [(11, "K"), (12, "L"), (13, "M"), (14, "N"), (15, "O"), (16, "P"),
                        (19, "S"), (20, "T"), (21, "U")]:
        _set(ws, total_row, col, f"=SUM({letter}{UTAMA_FIRST_ROW}:{letter}{UTAMA_LAST_ROW})", FORMULA_FILL)
    subtotal_utama_row = total_row + 1
    subtotal_pendukung_row = total_row + 2
    _set(ws, subtotal_utama_row, 1, "Subtotal - Kategori: Utama")
    _set(ws, subtotal_utama_row, 19, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Utama",$S${UTAMA_FIRST_ROW}:$S${UTAMA_LAST_ROW})', FORMULA_FILL)
    _set(ws, subtotal_utama_row, 20, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Utama",$T${UTAMA_FIRST_ROW}:$T${UTAMA_LAST_ROW})', FORMULA_FILL)
    _set(ws, subtotal_utama_row, 21, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Utama",$U${UTAMA_FIRST_ROW}:$U${UTAMA_LAST_ROW})', FORMULA_FILL)
    _set(ws, subtotal_pendukung_row, 1, "Subtotal - Kategori: Pendukung (dalam sheet ini)")
    _set(ws, subtotal_pendukung_row, 19, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Pendukung",$S${UTAMA_FIRST_ROW}:$S${UTAMA_LAST_ROW})', FORMULA_FILL)
    _set(ws, subtotal_pendukung_row, 20, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Pendukung",$T${UTAMA_FIRST_ROW}:$T${UTAMA_LAST_ROW})', FORMULA_FILL)
    _set(ws, subtotal_pendukung_row, 21, f'=SUMIF($E${UTAMA_FIRST_ROW}:$E${UTAMA_LAST_ROW},"Pendukung",$U${UTAMA_FIRST_ROW}:$U${UTAMA_LAST_ROW})', FORMULA_FILL)

    note_row = subtotal_pendukung_row + 2
    _set(ws, note_row, 1, "Keterangan & Rumus:")
    notes_utama = [
        "Sel KUNING = input manual. Sel ABU-ABU = formula (jangan diedit langsung).",
        "Setiap Produk dipecah 2 baris: Sistem = DB (database) dan Sistem = Engine (aplikasi/business logic).",
        "Required @ Current/Future = TPS x Rasio Resource per TPS.",
        "FINAL Required = Required@Future x (1+Buffer%) x HA Multiplier — angka untuk keputusan procurement/capacity planning.",
        "HA Multiplier = 1 (single node), 2 (active-passive/active-active 2 node), 3 (cluster 3 node).",
    ]
    for j, txt in enumerate(notes_utama, start=1):
        _subtitle(ws, note_row + j, f"• {txt}", 26)
    ws.freeze_panes = "F7"

    # =====================================================================
    # SHEET: Input - Infra Pendukung
    # =====================================================================
    ws2 = wb.create_sheet("Input - Infra Pendukung")
    _title(ws2, "Kalkulator Infrastruktur — INFRA PENDUKUNG (sizing manual, non-TPS)", 12)
    _subtitle(ws2, 2, "Contoh: Load Balancer, Monitoring Tools, Backup/DR, Firewall/Network Appliance, Jump Server, dsb. Sizing diisi manual.", 12)
    headers2 = ["No", "Komponen / Nama Infra", "Fungsi / Deskripsi", "Kategori\nInfrastruktur",
                "CPU\n(vCore)", "Memory\n(GB)", "Storage\n(GB)", "Buffer /\nHeadroom (%)", "HA\nMultiplier",
                "Final CPU\n(formula)", "Final Memory\n(formula)", "Final Storage\n(formula)"]
    _header_row(ws2, 4, headers2, col_widths=[5, 22, 30, 11, 9, 9, 9, 10, 8, 10, 10, 10])

    for i in range(n_pendukung):
        row = PENDUKUNG_FIRST_ROW + i
        data = pendukung[i] if i < len(pendukung) else {}
        _set(ws2, row, 1, i + 1)
        _set(ws2, row, 2, data.get("nama", ""), INPUT_FILL)
        _set(ws2, row, 3, data.get("fungsi", ""), INPUT_FILL)
        _set(ws2, row, 4, "Pendukung", INPUT_FILL)
        _set(ws2, row, 5, data.get("cpu"), INPUT_FILL)
        _set(ws2, row, 6, data.get("memory"), INPUT_FILL)
        _set(ws2, row, 7, data.get("storage"), INPUT_FILL)
        _set(ws2, row, 8, (data.get("buffer_pct", 20.0) / 100.0) if data else 0.2, INPUT_FILL)
        _set(ws2, row, 9, data.get("ha_multiplier", 1.0) if data else 1, INPUT_FILL)
        _set(ws2, row, 10, f"=E{row}*(1+H{row})*I{row}", FORMULA_FILL)
        _set(ws2, row, 11, f"=F{row}*(1+H{row})*I{row}", FORMULA_FILL)
        _set(ws2, row, 12, f"=G{row}*(1+H{row})*I{row}", FORMULA_FILL)

    total2_row = PENDUKUNG_LAST_ROW + 1
    _set(ws2, total2_row, 1, "TOTAL INFRA PENDUKUNG")
    for col, letter in [(5, "E"), (6, "F"), (7, "G"), (10, "J"), (11, "K"), (12, "L")]:
        _set(ws2, total2_row, col, f"=SUM({letter}{PENDUKUNG_FIRST_ROW}:{letter}{PENDUKUNG_LAST_ROW})", FORMULA_FILL)
    note2_row = total2_row + 2
    _set(ws2, note2_row, 1, "Keterangan:")
    notes_pendukung = [
        "Infra Pendukung disizing MANUAL langsung dalam satuan CPU/Memory/Storage (tidak diturunkan dari TPS).",
        "Final (CPU/Memory/Storage) = Base x (1+Buffer%) x HA Multiplier — total ini digabung dengan Infra Utama di sheet 'Output & Gap Analysis'.",
    ]
    for j, txt in enumerate(notes_pendukung, start=1):
        _subtitle(ws2, note2_row + j, f"• {txt}", 12)

    # =====================================================================
    # SHEET: Input - Storage Capacity
    # =====================================================================
    ws3 = wb.create_sheet("Input - Storage Capacity")
    _title(ws3, "Kalkulator Storage — basis Retensi & Housekeeping (per Platform > per Produk)", 16)
    _subtitle(ws3, 2, "Storage bersifat KUMULATIF — dihitung dari volume transaksi harian x ukuran data, dikali periode retensi.", 16)
    headers3 = ["No", "Platform", "Produk", "Volume Transaksi/Hari\n(Current, trx/hari)", "Volume Transaksi/Hari\n(Future, trx/hari)",
                "Ukuran Data\n/Transaksi (KB)", "Retensi\n(hari)", "Housekeeping/\nArchive Buffer (%)",
                "Storage Req\n@ Current (GB)", "Storage Req\n@ Future (GB)",
                "Buffer/\nHeadroom (%)", "Replication/HA\nMultiplier", "FINAL Storage Required (GB)",
                "Kuota Storage\nProduk (GB)", "Utilization vs\nKuota Produk (%)", "Status\n(vs Kuota Produk)"]
    _header_row(ws3, 4, headers3, col_widths=[5, 16, 22, 13, 13, 10, 8, 11, 11, 11, 9, 10, 11, 10, 11, 11])

    for i in range(n_storage):
        row = STORAGE_FIRST_ROW + i
        data = storage[i] if i < len(storage) else {}
        _set(ws3, row, 1, i + 1)
        _set(ws3, row, 2, data.get("platform", ""), INPUT_FILL)
        _set(ws3, row, 3, data.get("produk", ""), INPUT_FILL)
        _set(ws3, row, 4, data.get("vol_trx_current"), INPUT_FILL)
        _set(ws3, row, 5, data.get("vol_trx_future"), INPUT_FILL)
        _set(ws3, row, 6, data.get("ukuran_kb"), INPUT_FILL)
        _set(ws3, row, 7, data.get("retensi_hari"), INPUT_FILL)
        _set(ws3, row, 8, (data.get("housekeeping_pct", 15.0) / 100.0) if data else 0.15, INPUT_FILL)
        _set(ws3, row, 9, f'=IF(C{row}="","",D{row}*F{row}*G{row}/1048576*(1+H{row}))', FORMULA_FILL)
        _set(ws3, row, 10, f'=IF(C{row}="","",E{row}*F{row}*G{row}/1048576*(1+H{row}))', FORMULA_FILL)
        _set(ws3, row, 11, (data.get("buffer_pct", 20.0) / 100.0) if data else 0.2, INPUT_FILL)
        _set(ws3, row, 12, data.get("replication_multiplier", 1.0) if data else 1, INPUT_FILL)
        _set(ws3, row, 13, f'=IF(C{row}="","",J{row}*(1+K{row})*L{row})', FORMULA_FILL)
        _set(ws3, row, 14, data.get("kuota_storage"), INPUT_FILL)
        _set(ws3, row, 15, f'=IFERROR(M{row}/N{row},"")', FORMULA_FILL)
        _set(ws3, row, 16, f'=IF(N{row}="","",IF(O{row}>1,"OVER KUOTA",IF(O{row}>0.7,"WARNING","OK")))', FORMULA_FILL)

    total3_row = STORAGE_LAST_ROW + 1
    _set(ws3, total3_row, 1, "TOTAL STORAGE (semua produk)")
    for col, letter in [(9, "I"), (10, "J"), (13, "M"), (14, "N")]:
        _set(ws3, total3_row, col, f"=SUM({letter}{STORAGE_FIRST_ROW}:{letter}{STORAGE_LAST_ROW})", FORMULA_FILL)
    note3_row = total3_row + 2
    _set(ws3, note3_row, 1, "Keterangan & Rumus:")
    notes_storage = [
        "Storage dihitung KUMULATIF: Volume Transaksi/Hari x Ukuran Data/Transaksi x Retensi, dikonversi KB -> GB (dibagi 1.048.576).",
        "FINAL Storage Required = Required@Future x (1+Buffer%) x Replication/HA Multiplier.",
        "Kuota Storage Produk = ISI MANUAL, alokasi storage dari kapasitas platform induknya.",
    ]
    for j, txt in enumerate(notes_storage, start=1):
        _subtitle(ws3, note3_row + j, f"• {txt}", 16)

    # =====================================================================
    # SHEET: Input - Network Capacity
    # =====================================================================
    ws4 = wb.create_sheet("Input - Network Capacity")
    _title(ws4, "Kalkulator Network Capacity — basis TPS x Bandwidth/Transaksi, cascade ke Member/Link", 13)
    _subtitle(ws4, 2, "Bagian 1: total bandwidth required per Platform dari TPS. Bagian 2: cascade ke Member/Link sesuai % alokasi.", 13)
    _subtitle(ws4, 4, "1. Total Bandwidth Required per Platform (basis Transaksi)", 13)
    headers4 = ["No", "Platform", "Produk", "Current TPS", "Future TPS",
                "Bandwidth/Transaksi\n(KB) - Request", "Bandwidth/Transaksi\n(KB) - Response",
                "Overhead Protokol\n(%)", "Required @ Current\n(Mbps)", "Required @ Future\n(Mbps)",
                "Buffer/\nHeadroom (%)", "HA/Redundancy\nMultiplier", "FINAL Bandwidth\nRequired (Mbps)"]
    _header_row(ws4, 5, headers4, col_widths=[5, 16, 22, 10, 10, 12, 12, 10, 11, 11, 9, 10, 11])

    for i in range(n_network):
        row = NETWORK_FIRST_ROW + i
        data = network[i] if i < len(network) else {}
        _set(ws4, row, 1, i + 1)
        _set(ws4, row, 2, data.get("platform", ""), INPUT_FILL)
        _set(ws4, row, 3, data.get("produk", ""), INPUT_FILL)
        _set(ws4, row, 4, data.get("current_tps"), INPUT_FILL)
        _set(ws4, row, 5, data.get("future_tps"), INPUT_FILL)
        _set(ws4, row, 6, data.get("kb_req"), INPUT_FILL)
        _set(ws4, row, 7, data.get("kb_resp"), INPUT_FILL)
        _set(ws4, row, 8, (data.get("overhead_pct", 15.0) / 100.0) if data else 0.15, INPUT_FILL)
        _set(ws4, row, 9, f"=D{row}*(F{row}+G{row})*8*(1+H{row})/1000", FORMULA_FILL)
        _set(ws4, row, 10, f"=E{row}*(F{row}+G{row})*8*(1+H{row})/1000", FORMULA_FILL)
        _set(ws4, row, 11, (data.get("buffer_pct", 30.0) / 100.0) if data else 0.3, INPUT_FILL)
        _set(ws4, row, 12, data.get("ha_multiplier", 1.0) if data else 1, INPUT_FILL)
        _set(ws4, row, 13, f"=J{row}*(1+K{row})*L{row}", FORMULA_FILL)

    total4_row = NETWORK_LAST_ROW + 1
    _set(ws4, total4_row, 1, "TOTAL BANDWIDTH (semua produk)")
    for col, letter in [(9, "I"), (10, "J"), (13, "M")]:
        _set(ws4, total4_row, col, f"=SUM({letter}{NETWORK_FIRST_ROW}:{letter}{NETWORK_LAST_ROW})", FORMULA_FILL)

    _subtitle(ws4, MEMBER_FIRST_ROW - 2, "2. Cascade Total Bandwidth Platform ke Member/Link (distribusi trafik)", 13)
    headers4b = ["No", "Platform", "Member/Link\n(ISP/Switch/Node)", "% Alokasi\nTrafik",
                 "Bandwidth Required\n@ Member (Mbps)", "Existing Bandwidth\n@ Member (Mbps)",
                 "Sisa Kapasitas\n(Mbps)", "Utilization (%)", "Status"]
    _header_row(ws4, MEMBER_FIRST_ROW - 1, headers4b)

    for i in range(n_member):
        row = MEMBER_FIRST_ROW + i
        data = network_members[i] if i < len(network_members) else {}
        _set(ws4, row, 1, i + 1)
        _set(ws4, row, 2, data.get("platform", ""), INPUT_FILL)
        _set(ws4, row, 3, data.get("nama_link", ""), INPUT_FILL)
        _set(ws4, row, 4, (data.get("pct_alokasi", 0.0) / 100.0) if data else None, INPUT_FILL)
        _set(ws4, row, 5, f"=D{row}*SUMIF($B${NETWORK_FIRST_ROW}:$B${NETWORK_LAST_ROW},B{row},$M${NETWORK_FIRST_ROW}:$M${NETWORK_LAST_ROW})", FORMULA_FILL)
        _set(ws4, row, 6, data.get("existing_bandwidth"), INPUT_FILL)
        _set(ws4, row, 7, f'=IF(F{row}="","",F{row}-E{row})', FORMULA_FILL)
        _set(ws4, row, 8, f'=IFERROR(E{row}/F{row},"")', FORMULA_FILL)
        _set(ws4, row, 9, f'=IF(F{row}="","",IF(H{row}>1,"OVER CAPACITY",IF(H{row}>0.7,"WARNING","OK")))', FORMULA_FILL)

    note4_row = MEMBER_LAST_ROW + 2
    _set(ws4, note4_row, 1, "Keterangan & Rumus:")
    notes_network = [
        "Required Bandwidth (Mbps) = TPS x (KB Request + KB Response) x 8 (byte->bit) x (1+Overhead%) / 1000.",
        "FINAL Bandwidth Required = Required@Future x (1+Buffer%) x HA/Redundancy Multiplier.",
        "Bagian 2: total FINAL Bandwidth per Platform di-SUMIF lalu didistribusikan ke Member/Link sesuai % Alokasi Trafik.",
        "Status: OK (<=70%), WARNING (70-100%), OVER CAPACITY (>100%).",
    ]
    for j, txt in enumerate(notes_network, start=1):
        _subtitle(ws4, note4_row + j, f"• {txt}", 13)

    # =====================================================================
    # SHEET: Dashboard
    # =====================================================================
    wsd = wb.create_sheet("Dashboard", 0)
    _title(wsd, "DASHBOARD KAPASITAS INFRASTRUKTUR", 11)
    _subtitle(wsd, 2, "Layer 1: Platform (kapasitas existing/shared pool) — Layer 2: Produk (kuota alokasi dalam platform). "
                      "Hijau <=70% | Kuning 70-90% | Merah >90%.", 11)
    _subtitle(wsd, 4, "LAYER 1 — PER PLATFORM (vs Kapasitas Existing Shared Pool)", 11)
    headers_l1 = ["Platform", "CPU Required\n(vCore)", "CPU Existing\n(vCore)", "CPU Util%", "Status CPU",
                  "Memory Required\n(GB)", "Memory Existing\n(GB)", "Memory Util%", "Status Memory",
                  "Storage Required\n(GB)", "Network Required\n(Mbps)"]
    _header_row(wsd, 5, headers_l1, col_widths=[18, 12, 12, 9, 10, 12, 12, 9, 10, 12, 13])

    L1_FIRST = 6
    L1_LAST = L1_FIRST + n_platform - 1
    for i in range(n_platform):
        row = L1_FIRST + i
        p = platforms[i] if i < len(platforms) else {}
        _set(wsd, row, 1, p.get("platform", ""), INPUT_FILL)
        _set(wsd, row, 2, f'=IF(A{row}="","",SUMIF(\'Input - Infra Utama\'!$B${UTAMA_FIRST_ROW}:$B${UTAMA_LAST_ROW},A{row},\'Input - Infra Utama\'!$S${UTAMA_FIRST_ROW}:$S${UTAMA_LAST_ROW})+IF(A{row}="","",SUMIF(\'Input - Infra Pendukung\'!$B${PENDUKUNG_FIRST_ROW}:$B${PENDUKUNG_LAST_ROW},A{row},\'Input - Infra Pendukung\'!$J${PENDUKUNG_FIRST_ROW}:$J${PENDUKUNG_LAST_ROW})))', FORMULA_FILL)
        _set(wsd, row, 3, p.get("cpu"), INPUT_FILL)
        _set(wsd, row, 4, f'=IFERROR(B{row}/C{row},"")', FORMULA_FILL)
        _set(wsd, row, 5, f'=IF(D{row}="","",IF(D{row}>0.9,"MERAH",IF(D{row}>0.7,"KUNING","HIJAU")))', FORMULA_FILL)
        _set(wsd, row, 6, f'=IF(A{row}="","",SUMIF(\'Input - Infra Utama\'!$B${UTAMA_FIRST_ROW}:$B${UTAMA_LAST_ROW},A{row},\'Input - Infra Utama\'!$T${UTAMA_FIRST_ROW}:$T${UTAMA_LAST_ROW})+SUMIF(\'Input - Infra Pendukung\'!$B${PENDUKUNG_FIRST_ROW}:$B${PENDUKUNG_LAST_ROW},A{row},\'Input - Infra Pendukung\'!$K${PENDUKUNG_FIRST_ROW}:$K${PENDUKUNG_LAST_ROW}))', FORMULA_FILL)
        _set(wsd, row, 7, p.get("memory"), INPUT_FILL)
        _set(wsd, row, 8, f'=IFERROR(F{row}/G{row},"")', FORMULA_FILL)
        _set(wsd, row, 9, f'=IF(H{row}="","",IF(H{row}>0.9,"MERAH",IF(H{row}>0.7,"KUNING","HIJAU")))', FORMULA_FILL)
        _set(wsd, row, 10, f'=IF(A{row}="","",SUMIF(\'Input - Storage Capacity\'!$B${STORAGE_FIRST_ROW}:$B${STORAGE_LAST_ROW},A{row},\'Input - Storage Capacity\'!$M${STORAGE_FIRST_ROW}:$M${STORAGE_LAST_ROW})+SUMIF(\'Input - Infra Pendukung\'!$B${PENDUKUNG_FIRST_ROW}:$B${PENDUKUNG_LAST_ROW},A{row},\'Input - Infra Pendukung\'!$L${PENDUKUNG_FIRST_ROW}:$L${PENDUKUNG_LAST_ROW}))', FORMULA_FILL)
        _set(wsd, row, 11, f'=IF(A{row}="","",SUMIF(\'Input - Network Capacity\'!$B${NETWORK_FIRST_ROW}:$B${NETWORK_LAST_ROW},A{row},\'Input - Network Capacity\'!$M${NETWORK_FIRST_ROW}:$M${NETWORK_LAST_ROW}))', FORMULA_FILL)

    l2_title_row = L1_LAST + 3
    _subtitle(wsd, l2_title_row, "LAYER 2 — PER PRODUK (vs Kuota Alokasi dalam Platform)", 11)
    headers_l2 = ["Platform", "Produk", "CPU Required\n(vCore)", "CPU Kuota\n(vCore)", "CPU Util%", "Status CPU",
                  "Memory Required\n(GB)", "Memory Kuota\n(GB)", "Memory Util%", "Status Memory",
                  "Storage Status\n(dari sheet Storage)"]
    _header_row(wsd, l2_title_row + 1, headers_l2)

    unique_produk: List[tuple] = []
    seen = set()
    for r in infra_utama:
        key = (r.get("platform"), r.get("produk"))
        if key not in seen and key[1]:
            seen.add(key)
            unique_produk.append(key)

    L2_FIRST = l2_title_row + 2
    n_l2 = max(MIN_PRODUK_ROWS, len(unique_produk))
    L2_LAST = L2_FIRST + n_l2 - 1
    for i in range(n_l2):
        row = L2_FIRST + i
        if i < len(unique_produk):
            plat, produk = unique_produk[i]
        else:
            plat, produk = "", ""
        _set(wsd, row, 1, plat, INPUT_FILL)
        _set(wsd, row, 2, produk, INPUT_FILL)
        _set(wsd, row, 3, f'=IF(B{row}="","",SUMIF(\'Input - Infra Utama\'!$C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},B{row},\'Input - Infra Utama\'!$S${UTAMA_FIRST_ROW}:$S${UTAMA_LAST_ROW}))', FORMULA_FILL)
        _set(wsd, row, 4, f'=IF(B{row}="","",IFERROR(INDEX(\'Input - Infra Utama\'!$V${UTAMA_FIRST_ROW}:$V${UTAMA_LAST_ROW},MATCH(B{row},\'Input - Infra Utama\'!$C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},0)),""))', FORMULA_FILL)
        _set(wsd, row, 5, f'=IFERROR(C{row}/D{row},"")', FORMULA_FILL)
        _set(wsd, row, 6, f'=IF(E{row}="","",IF(E{row}>0.9,"MERAH",IF(E{row}>0.7,"KUNING","HIJAU")))', FORMULA_FILL)
        _set(wsd, row, 7, f'=IF(B{row}="","",SUMIF(\'Input - Infra Utama\'!$C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},B{row},\'Input - Infra Utama\'!$T${UTAMA_FIRST_ROW}:$T${UTAMA_LAST_ROW}))', FORMULA_FILL)
        _set(wsd, row, 8, f'=IF(B{row}="","",IFERROR(INDEX(\'Input - Infra Utama\'!$W${UTAMA_FIRST_ROW}:$W${UTAMA_LAST_ROW},MATCH(B{row},\'Input - Infra Utama\'!$C${UTAMA_FIRST_ROW}:$C${UTAMA_LAST_ROW},0)),""))', FORMULA_FILL)
        _set(wsd, row, 9, f'=IFERROR(G{row}/H{row},"")', FORMULA_FILL)
        _set(wsd, row, 10, f'=IF(I{row}="","",IF(I{row}>0.9,"MERAH",IF(I{row}>0.7,"KUNING","HIJAU")))', FORMULA_FILL)
        _set(wsd, row, 11, f'=IFERROR(INDEX(\'Input - Storage Capacity\'!$P${STORAGE_FIRST_ROW}:$P${STORAGE_LAST_ROW},MATCH(B{row},\'Input - Storage Capacity\'!$C${STORAGE_FIRST_ROW}:$C${STORAGE_LAST_ROW},0)),"")', FORMULA_FILL)

    note_d_row = L2_LAST + 2
    _set(wsd, note_d_row, 1, "Keterangan:")
    notes_dash = [
        "LAYER 1 (Platform): Required = agregat semua produk dalam platform. Existing = kapasitas fisik shared pool — ISI MANUAL (kolom C & G).",
        "LAYER 2 (Produk): Required = rollup DB+Engine untuk 1 produk. Kuota = alokasi dari jatah platform-nya — ISI MANUAL di sheet 'Input - Infra Utama' (kolom V, W).",
        "Sebuah Produk bisa HIJAU di Layer 2 tapi Platform-nya MERAH di Layer 1 (produk lain di platform sama menghabiskan kapasitas) — cek kedua layer.",
        "Warna: HIJAU <=70% (cukup), KUNING 70-90% (mulai penuh), MERAH >90% (kritis/perlu ekspansi).",
    ]
    for j, txt in enumerate(notes_dash, start=1):
        _subtitle(wsd, note_d_row + j, f"• {txt}", 11)

    # =====================================================================
    # SHEET: Output & Gap Analysis
    # =====================================================================
    wso = wb.create_sheet("Output & Gap Analysis")
    _title(wso, "Output: Ringkasan Kebutuhan Infrastruktur vs Kapasitas Existing", 6)

    _subtitle(wso, 3, "A. Total FINAL Required (CPU/Memory dari Infra Utama+Pendukung; Storage dari sheet Storage Capacity), per Kategori", 6)
    _header_row(wso, 4, ["Kategori", "CPU (vCore)", "Memory (GB)", "Storage (GB)"])
    _set(wso, 5, 1, "Infra Utama - Kategori 'Utama'")
    _set(wso, 5, 2, f"='Input - Infra Utama'!S{subtotal_utama_row}", FORMULA_FILL)
    _set(wso, 5, 3, f"='Input - Infra Utama'!T{subtotal_utama_row}", FORMULA_FILL)
    _set(wso, 5, 4, 0)
    _set(wso, 6, 1, "Infra Utama - Kategori 'Pendukung'")
    _set(wso, 6, 2, f"='Input - Infra Utama'!S{subtotal_pendukung_row}", FORMULA_FILL)
    _set(wso, 6, 3, f"='Input - Infra Utama'!T{subtotal_pendukung_row}", FORMULA_FILL)
    _set(wso, 6, 4, 0)
    _set(wso, 7, 1, "Infra Pendukung (sheet tersendiri)")
    _set(wso, 7, 2, f"='Input - Infra Pendukung'!J{total2_row}", FORMULA_FILL)
    _set(wso, 7, 3, f"='Input - Infra Pendukung'!K{total2_row}", FORMULA_FILL)
    _set(wso, 7, 4, f"='Input - Infra Pendukung'!L{total2_row}", FORMULA_FILL)
    _set(wso, 8, 1, "Storage Capacity (sheet tersendiri, basis retensi)")
    _set(wso, 8, 2, 0)
    _set(wso, 8, 3, 0)
    _set(wso, 8, 4, f"='Input - Storage Capacity'!M{total3_row}", FORMULA_FILL)
    _set(wso, 9, 1, "TOTAL KESELURUHAN")
    _set(wso, 9, 2, "=SUM(B5:B8)", FORMULA_FILL)
    _set(wso, 9, 3, "=SUM(C5:C8)", FORMULA_FILL)
    _set(wso, 9, 4, "=SUM(D5:D8)", FORMULA_FILL)

    _subtitle(wso, 11, "B. Kapasitas Existing / Tersedia (Shared Pool) — ISI MANUAL", 6)
    _header_row(wso, 12, ["Resource", "Kapasitas Existing", "Satuan"])
    total_cpu = sum(float(p.get("cpu", 0) or 0) for p in platforms)
    total_mem = sum(float(p.get("memory", 0) or 0) for p in platforms)
    total_storage_cap = sum(float(p.get("storage", 0) or 0) for p in platforms)
    total_network_cap = sum(float(p.get("network", 0) or 0) for p in platforms)
    _set(wso, 13, 1, "CPU"); _set(wso, 13, 2, total_cpu, INPUT_FILL); _set(wso, 13, 3, "vCore")
    _set(wso, 14, 1, "Memory"); _set(wso, 14, 2, total_mem, INPUT_FILL); _set(wso, 14, 3, "GB")
    _set(wso, 15, 1, "Storage"); _set(wso, 15, 2, total_storage_cap, INPUT_FILL); _set(wso, 15, 3, "GB")
    _set(wso, 16, 1, "Network"); _set(wso, 16, 2, total_network_cap, INPUT_FILL); _set(wso, 16, 3, "Mbps")

    _subtitle(wso, 18, "C. Analisis Gap & Utilization (Total Keseluruhan vs Existing)", 6)
    _header_row(wso, 19, ["Resource", "Total Required\n(Final)", "Existing Capacity",
                          "Sisa Kapasitas\n(Existing-Required)", "Utilization (%)", "Status"])
    gap_rows = [
        ("CPU", "=B9", "=B13"),
        ("Memory", "=C9", "=B14"),
        ("Storage", "=D9", "=B15"),
        ("Network", f"=SUM('Input - Network Capacity'!M{NETWORK_FIRST_ROW}:M{NETWORK_LAST_ROW})", "=B16"),
    ]
    for i, (label, req_formula, exist_formula) in enumerate(gap_rows):
        row = 20 + i
        _set(wso, row, 1, label)
        _set(wso, row, 2, req_formula, FORMULA_FILL)
        _set(wso, row, 3, exist_formula, FORMULA_FILL)
        _set(wso, row, 4, f"=C{row}-B{row}", FORMULA_FILL)
        _set(wso, row, 5, f"=IFERROR(B{row}/C{row},0)", FORMULA_FILL)
        _set(wso, row, 6, f'=IF(C{row}=0,"N/A",IF(E{row}>1,"OVER CAPACITY",IF(E{row}>0.8,"WARNING - Perlu Review","OK")))', FORMULA_FILL)

    note_o_row = 25
    _set(wso, note_o_row, 1, "Keterangan:")
    notes_output = [
        "Section A: CPU & Memory dari sheet 'Input - Infra Utama' + 'Input - Infra Pendukung'. Storage dari sheet 'Input - Storage Capacity'.",
        "Section B (kuning) diisi MANUAL — kapasitas existing/tersedia infrastruktur saat ini, termasuk Network.",
        "Section C = perbandingan TOTAL KESELURUHAN vs kapasitas existing. Breakdown per platform/produk ada di sheet 'Dashboard'.",
        "Status: OK (<=80%), WARNING - Perlu Review (80-100%), OVER CAPACITY (>100%).",
    ]
    for j, txt in enumerate(notes_output, start=1):
        _subtitle(wso, note_o_row + j, f"• {txt}", 6)

    for ws_ in wb.worksheets:
        ws_.sheet_view.showGridLines = False

    return wb


def build_template_excel_bytes(DATA: Dict[str, Any]) -> bytes:
    import io
    wb = build_template_workbook(DATA)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
